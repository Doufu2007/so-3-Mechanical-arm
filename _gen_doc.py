# -*- coding: utf-8 -*-
"""生成《SE(3) 等变逆动力学：基座姿态零样本泛化》论文方向与实施方案 Word 文档。"""
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn

EAST = '宋体'
HEAD_EAST = '黑体'
ASCII = 'Calibri'

doc = Document()

# 页面默认样式
normal = doc.styles['Normal']
normal.font.name = ASCII
normal.font.size = Pt(11)
normal._element.rPr.rFonts.set(qn('w:eastAsia'), EAST)


def _set(run, size=11, bold=False, east=EAST, ascii_=ASCII, color=None, italic=False):
    run.font.name = ascii_
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = rpr.makeelement(qn('w:rFonts'), {})
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), east)
    return run


def h1(text):
    p = doc.add_paragraph()
    _set(p.add_run(text), size=15, bold=True, east=HEAD_EAST, color=(0x1F, 0x3B, 0x63))
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(6)
    return p


def h2(text):
    p = doc.add_paragraph()
    _set(p.add_run(text), size=12.5, bold=True, east=HEAD_EAST, color=(0x2E, 0x54, 0x86))
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    return p


def body(text, bold=False, size=11):
    p = doc.add_paragraph()
    _set(p.add_run(text), size=size, bold=bold)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.25
    return p


def bullet(text, bold_prefix=None):
    p = doc.add_paragraph(style='List Bullet')
    if bold_prefix:
        _set(p.add_run(bold_prefix), bold=True)
    _set(p.add_run(text), size=11)
    p.paragraph_format.space_after = Pt(2)
    return p


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        _set(hdr[i].paragraphs[0].add_run(h), size=10.5, bold=True, east=HEAD_EAST)
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = ''
            _set(cells[i].paragraphs[0].add_run(v), size=10.5)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = w
    doc.add_paragraph()
    return t


# ======================= 封面标题 =======================
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set(title.add_run('SE(3) 等变神经逆动力学'), size=20, bold=True, east=HEAD_EAST, color=(0x14, 0x2D, 0x4C))
sub = doc.add_paragraph()
sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set(sub.add_run('—— 基座安装姿态的零样本泛化 ——'), size=14, bold=True, east=HEAD_EAST, color=(0x3B, 0x5A, 0x88))
meta = doc.add_paragraph()
meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
_set(meta.add_run('论文方向与实施方案  ·  v1  ·  2026-09-03'), size=10.5, color=(0x80, 0x80, 0x80))

# ======================= 0. 一页定论 =======================
h1('0. 一页定论')
body('这篇论文做的是「SE(3) 等变逆动力学」：利用“重力项对重力方向线性”这一物理结构，'
     '让同一个神经网络模型训练一次，机械臂换任何安装姿态（正装 / 倒装 / 壁挂）都零重训直接泛化。')
p = doc.add_paragraph()
p.paragraph_format.space_before = Pt(6)
_set(p.add_run('一句话命题：'), size=11, bold=True)
_set(p.add_run('不是发明新网络，而是把教科书已知的“重力对重力方向线性”做进学习模型的结构约束，'
               '再用“换姿态零重训”这个没人做过的实验证明这个约束值钱——黑箱做不到，结构做得到。'), size=11, bold=True)

# ======================= 1. 命题链 =======================
h1('1. 命题链')
table(
    ['环节', '内容'],
    [
        ['问题', '现有神经逆动力学把基座姿态写死（重力方向是写死的常量）。机械臂一换安装姿态——移动底盘、'
                 '多工位、倒装壁挂——就必须重训。但真机场景里换姿态是常态。'],
        ['方法', '用 SE(3) 物理结构：惯性项 H(q)、科氏项 c(q,q̇) 对基座旋转不变，只有重力项 g(q) 随重力方向 '
                 'g_vec 线性变化。把 g_vec 从“写死的常量”升级成“显式输入”，模型就天然对基座姿态 SO(3) 等变。'
                 '不需要任何新网络架构——现有代码 V = −m(g_vec·p_com) 已经是这个性质，g_vec 现在是 buffer，'
                 '改成变量就成。'],
        ['证据', '生死门三组对照（见 §3）。'],
        ['卖点', '“换安装姿态零样本泛化”是真机（移动底盘 / 多工位）的真实需求，不是又一轮旧方法拼装。'],
    ],
)

# ======================= 2. 为什么新 =======================
h1('2. 为什么它是“新”的（对审稿人的那句话）')
body('不是发明新网络，而是把教科书已知的“重力对重力方向线性”做进学习模型的结构约束，'
     '再用“换姿态零重训”这个没人做过的实验证明这个约束值钱——黑箱做不到，结构做得到。')

h2('2.1 空白判定：高置信，但非铁证')
bullet('做的是前向 Port-Hamiltonian 神经 ODE + 自由浮体 / 飞行器，无机械臂关节空间、无逆动力学辨识、无基座姿态。',
       bold_prefix='Duong 等（T-RO 2024）：')
bullet('定向检索 “equivariant inverse dynamics + manipulator + base orientation” 与 '
       '“gravity reorientation + zero retraining” 均零命中。', bold_prefix='检索结论：')
bullet('① 搜索引擎看不到的（刚投 preprint、闭门工作、学位论文）；② 范式级已被 Duong / Seo 占据，增量是'
       '“应用级组合”；③ “重力线性于重力方向”是教科书，物理上不新鲜——新鲜的是“做成学习模型的结构约束 + '
       '黑箱对照 + 零重训证据”。', bold_prefix='三层边界：')
bullet('低但非零。定稿前必补：Google Scholar “equivariant inverse dynamics” 近两年 + Duong 在投 preprint 扫描。',
       bold_prefix='残余风险：')

# ======================= 3. 生死门 =======================
h1('3. 生死门实验设计（全篇命门）')
body('核心判断：结构模型的“等变”是数学保证（必然成立），所以卖点不在“结构换姿态后不变”，'
     '而在“黑箱换姿态后是否显著变差”。')
h2('3.1 三组对照（换姿态后全部零重训）')
table(
    ['组', '模型输入', '换姿态后（零重训）预期'],
    [
        ['结构模型（g 对 g_vec 线性）', 'q, q̇, q̈ + 重力方向 g_vec（显式变量）', '保持'],
        ['裸黑箱 MLP', 'q, q̇, q̈（不含重力方向）', '失效'],
        ['带重力方向输入的黑箱 MLP', 'q, q̇, q̈ + g_vec', '？ ← 真正的对手'],
    ],
)
body('第三组是生死门：若“带重力方向输入的黑箱”零重训后也基本泛化（黑箱有能力学出重力方向特征），'
     '则结构等变无独特优势，卖点塌；若黑箱即使喂了重力方向、零重训仍明显变差（它没被约束“g 对 g_vec 必须线性”，'
     '学的是非线性近似），则结构约束值钱。')
body('这个第三组对照是 novelty 能否站住的唯一一根柱子，先想清楚再动手。', bold=True)

# ======================= 4. 技术路线 =======================
h1('4. 技术路线')
h2('4.1 代码改动（增量小）')
bullet('把 se3_physics.py 里 g_vec 从 register_buffer 改为前向输入（可变变量）。')
bullet('重力项保持 V = −m(g_vec·p_com) 线性形式，g_vec 由外部注入。')
bullet('数据管线增加“姿态参数化”：用一组基座旋转 R ∈ SO(3) 生成多姿态数据。')
h2('4.2 数据生成（复用现有 MuJoCo 管线）')
bullet('用现有 franka MuJoCo 环境，绕竖直轴 / 水平轴生成多个基座安装姿态。')
bullet('每个姿态生成 (q, q̇, q̈, τ) 数据集；训练只用“正装”一个姿态，测试用其余姿态（倒装 / 壁挂 / 侧装）。')
h2('4.3 真机验证（Bobac3）')
bullet('真机不需要“可重装姿态”——只做两件事：① 常规闭环（重力补偿轨迹跟踪）验证模型在真实机器人上闭环可行；'
       '② 线性性验证：在真机上小角度改变基座安装角，验证重力项随重力方向线性（等价于验证 g(q) 对 g_vec 线性）。')

# ======================= 5. 与 A 的关系 + 保底 =======================
h1('5. 与既有工程类版（A）的关系 + 保底策略')
bullet('A（重参数化工程类版）已写好 manuscript.md，作为三区保底底稿保留，不丢弃。', bold_prefix='底稿：')
bullet('B（等变逆动力学）是主攻方向，冲二区。', bold_prefix='主攻：')
bullet('B 冲二区（有真实机会）+ A 三区落袋保底 48 + 96 是“冲”出来的，不裸奔。', bold_prefix='策略：')

# ======================= 6. 风险清单 =======================
h1('6. 风险清单')
table(
    ['#', '风险', '后果 / 应对'],
    [
        ['1', '生死门第三组跑不出“结构显著赢黑箱”', '卖点塌，退三区（最大风险）'],
        ['2', '审稿人“重力线性是教科书，trivial”', '靠黑箱对照正面接住；接不住则 novelty 被否'],
        ['3', '空白判定非铁证', '定稿前补文献扫描（Google Scholar + Duong preprint）'],
        ['4', '时间被低估', '实际 2–4 周（改代码 + 多姿态数据 + 三组模型 + 测试），非 1–2 周'],
    ],
)

# ======================= 7. 时间线 =======================
h1('7. 时间线（草案）')
table(
    ['阶段', '内容'],
    [
        ['第 1 周', '改代码（g_vec 变变量）+ 生成多姿态数据 + 训练结构模型'],
        ['第 2 周', '训练两组黑箱对照 + 换姿态零重训测试（生死门出结果）'],
        ['第 3 周', '视生死门结果决定继续 / 调整 + 文献补扫描'],
        ['第 4 周', '写 manuscript + 真机闭环 / 线性性验证'],
    ],
)

out = 'SE3等变逆动力学_论文方向与实施方案.docx'
doc.save(out)
print('saved:', out)
