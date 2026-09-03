# -*- coding: utf-8 -*-
"""生成《SO(3) 项目 vs SE(3) 项目对比》Word 文档到桌面"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

doc = Document()

# ---------- 中文字体设置 ----------
def set_cn(run, size=None, bold=None, color=None):
    run.font.name = 'Calibri'
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = rPr.makeelement(qn('w:rFonts'), {})
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), '微软雅黑')
    if size: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    if color: run.font.color.rgb = color

normal = doc.styles['Normal']
normal.font.name = 'Calibri'
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), '微软雅黑')

def P(text, bold=False, size=10.5, color=None, space_after=6, align=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn(run, size=size, bold=bold, color=color)
    p.paragraph_format.space_after = Pt(space_after)
    if align: p.alignment = align
    return p

# ===================== 封面标题 =====================
title = doc.add_heading(level=0)
r = title.add_run('原来的 SO(3) 项目  vs  现在的 SE(3) 项目')
set_cn(r, size=20, bold=True)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER

# 一句话
one = doc.add_paragraph()
run = one.add_run('一句话：李群放的地方换了——从「输入编码」搬到了「动力学结构」里。')
set_cn(run, size=13, bold=True, color=RGBColor(0xC0, 0x00, 0x00))
one.alignment = WD_ALIGN_PARAGRAPH.CENTER
one.paragraph_format.space_after = Pt(14)

# ===================== 对比表 =====================
rows = [
    ['对比维度', '原来的 SO(3) 项目', '现在的 SE(3) 项目'],
    ['李群在哪',
     '放在输入编码：q → sin/cos 或旋转矩阵',
     '放在动力学结构：M(q)=Σ Jᵀ G J（SE(3) 空间雅可比 × 学的空间惯量）'],
    ['那个李群是真的吗',
     '基本是假的：SO3LieGroup 是死代码（从没进前向）；真正用的只是 sin(q)/cos(q) 拼接；真旋转矩阵版本从没训练',
     '真的：惯性矩阵的 SPD 是几何保证的（不是 M·Mᵀ 后处理），只学 80 个物理惯量参数'],
    ['核心主张',
     '「SO(3) 编码避开奇异点」',
     '「SE(3) 结构让模型学到可辨识的 base 参数」'],
    ['结果诚实度',
     '奇异点证据是编的（singularity_analysis.py 里「欧拉角」那条曲线是手调的解析公式，不是测的）；SO(3) 编码实测 ≈ 没帮助',
     '代码重写成 paper/（权威版），修了旧实现 3 处会改结论的错，负结果也老实写'],
    ['模型',
     'DeLaN-FFNN：任意神经网络 H(q) + 残差分支',
     '结构化惯量（80 参数）+ 残差分支'],
]

t = doc.add_table(rows=len(rows), cols=len(rows[0]))
t.style = 'Light Grid Accent 1'
t.alignment = WD_TABLE_ALIGNMENT.CENTER

for i, row in enumerate(rows):
    for j, cell in enumerate(row):
        c = t.cell(i, j)
        c.text = ''
        para = c.paragraphs[0]
        run = para.add_run(str(cell))
        if i == 0:
            set_cn(run, size=10.5, bold=True)
            shd = c._element.get_or_add_tcPr().makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'D9E2F3'})
            c._element.get_or_add_tcPr().append(shd)
        elif j == 0:
            set_cn(run, size=10, bold=True)
            shd = c._element.get_or_add_tcPr().makeelement(qn('w:shd'), {qn('w:val'): 'clear', qn('w:fill'): 'F2F2F2'})
            c._element.get_or_add_tcPr().append(shd)
        else:
            set_cn(run, size=10)

# 列宽：首列窄，后两列宽
from docx.shared import Cm
widths = [Cm(2.8), Cm(6.6), Cm(6.6)]
for j, w in enumerate(widths):
    for i in range(len(rows)):
        t.cell(i, j).width = w

out = os.path.join(os.path.expanduser('~'), 'Desktop', 'SO3_vs_SE3_项目对比.docx')
doc.save(out)
print('DONE:', out)
